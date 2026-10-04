"""Выгрузка записей в Excel с теми же фильтрами, что и в списках."""

from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.models import Building, Responsible, Room

THIN = Side(style='thin', color='C8D1DE')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='0B5CAD')
HEADER_FONT = Font(bold=True, color='FFFFFF')
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color='4D5B6C')
INVALID_FILL = PatternFill('solid', fgColor='FDEEEE')
INVALID_FONT = Font(bold=True, color='C62F2F')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
CELL_ALIGN = Alignment(vertical='top', wrap_text=True)
HEADER_ROW = 5

EXPECTED_WAVEFORM = 'синусоидальная'


def _num(value):
    return None if value is None or value == '' else float(value)


def _fmt_date(value):
    return value.strftime('%d.%m.%Y') if value else ''


def _fmt_iso(value):
    try:
        return datetime.strptime(value, '%Y-%m-%d').strftime('%d.%m.%Y')
    except (TypeError, ValueError):
        return None


def _person(person):
    if not person:
        return ''
    return ' '.join(part for part in (person.last_name, person.first_name, person.patronymic) if part)


def _instruments(record):
    lines = []
    for instrument in record.measurement_instruments.all():
        valid_until = _fmt_date(instrument.next_calibration_date) or 'н/д'
        lines.append(f'{instrument.name} {instrument.type} (№ {instrument.serial_number}), годен до {valid_until}')
    return '\n'.join(lines)


def _out_of_range(value, low, high):
    value, low, high = _num(value), _num(low), _num(high)
    if value is None:
        return False
    return (low is not None and value < low) or (high is not None and value > high)


def _ids(params, key):
    return [int(part) for part in str(params.get(key) or '').split(',') if part.strip().isdigit()]


def _period_text(params):
    if params.get('date'):
        return f"за {_fmt_iso(params.get('date')) or params.get('date')}"
    start, end = _fmt_iso(params.get('start_date')), _fmt_iso(params.get('end_date'))
    if start and end:
        return f'с {start} по {end}'
    if start:
        return f'с {start}'
    if end:
        return f'по {end}'
    return 'за всё время'


def _filters_summary(params, scope):
    parts = [f'Период: {_period_text(params)}']
    if params.get('mine') in ('1', 'true'):
        parts.append('только мои здания' if scope == 'buildings' else 'только мои помещения')

    responsible_ids = _ids(params, 'responsible')
    if responsible_ids:
        names = [_person(p) for p in Responsible.objects.filter(id__in=responsible_ids).order_by('last_name')]
        parts.append('Ответственные: ' + ', '.join(names))

    if scope == 'buildings':
        building_ids = _ids(params, 'buildings')
        if building_ids:
            numbers = Building.objects.filter(id__in=building_ids).values_list('building_number', flat=True)
            parts.append('Здания: ' + ', '.join(sorted(numbers)))
    else:
        room_ids = _ids(params, 'rooms')
        if room_ids:
            rooms = Room.objects.filter(id__in=room_ids).select_related('building').order_by('building__building_number', 'room_number')
            parts.append('Помещения: ' + ', '.join(
                f'{r.room_number} (зд. {r.building.building_number})' if r.building_id else r.room_number for r in rooms
            ))
    return ' · '.join(parts)


def _filename(title, params):
    start, end = _fmt_iso(params.get('start_date')), _fmt_iso(params.get('end_date'))
    if start or end:
        suffix = f"{start or '…'}–{end or '…'}"
    elif params.get('date'):
        suffix = _fmt_iso(params.get('date')) or params.get('date')
    else:
        suffix = f"от {timezone.localdate().strftime('%d.%m.%Y')}"
    return f'{title} {suffix}.xlsx'


def _build_workbook(sheet_title, title, summary, columns, rows):
    """columns: [(заголовок, ширина, формат числа)]; rows: [[(значение, вне нормы)]]"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws['A1'] = title
    ws['A1'].font = TITLE_FONT
    ws['A2'] = summary
    ws['A2'].font = NOTE_FONT
    ws['A3'] = f"Сформировано {timezone.localtime().strftime('%d.%m.%Y %H:%M')} · строк: {len(rows)}"
    ws['A3'].font = NOTE_FONT

    for col, (header, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[HEADER_ROW].height = 32

    for row_index, row in enumerate(rows, start=HEADER_ROW + 1):
        for col, (value, invalid) in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=col, value=value)
            cell.border = BORDER
            cell.alignment = CELL_ALIGN
            number_format = columns[col - 1][2]
            if number_format and isinstance(value, float):
                cell.number_format = number_format
            if invalid:
                cell.fill = INVALID_FILL
                cell.font = INVALID_FONT

    last_col = get_column_letter(len(columns))
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.auto_filter.ref = f'A{HEADER_ROW}:{last_col}{HEADER_ROW + max(len(rows), 1)}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{HEADER_ROW}:{HEADER_ROW}'
    return wb


def _response(wb, filename):
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
    response['Access-Control-Expose-Headers'] = 'Content-Disposition'
    return response


ROOM_COLUMNS = [
    ('Дата', 12, None),
    ('Здание', 9, None),
    ('Помещение', 12, None),
    ('Тип помещения', 15, None),
    ('Ответственный', 26, None),
    ('Время', 8, None),
    ('Температура, °C', 13, '0.00'),
    ('Влажность, %', 12, '0.0'),
    ('Давление, кПа', 12, '0.00'),
    ('Давление, мм рт. ст.', 13, '0.00'),
    ('Напряжение, В', 12, '0.00'),
    ('Частота, Гц', 11, '0.00'),
    ('Радиационный фон, мкЗв/ч', 14, '0.00'),
    ('Отклонения от нормы', 26, None),
    ('Средства измерений', 48, None),
]


def _room_type(room):
    if room.is_storage:
        return 'КВХ'
    if room.has_additional_parameters:
        return 'С доп. параметрами'
    return 'Обычное'


def _room_metrics(room):
    extra = getattr(room, 'additional_parameters', None) if room.has_additional_parameters else None
    return [
        ('temperature_celsius', 'температура', room.temperature_min, room.temperature_max),
        ('humidity_percentage', 'влажность', room.humidity_min, room.humidity_max),
        ('pressure_kpa', 'давление (кПа)', room.pressure_min_kpa, room.pressure_max_kpa),
        ('pressure_mmhg', 'давление (мм рт. ст.)', room.pressure_min_mmhg, room.pressure_max_mmhg),
        ('voltage', 'напряжение', extra.voltage_min if extra else None, extra.voltage_max if extra else None),
        ('frequency', 'частота', extra.frequency_min if extra else None, extra.frequency_max if extra else None),
        ('radiation', 'радиационный фон', extra.radiation_min if extra else None, extra.radiation_max if extra else None),
    ]


def export_rooms(records, params):
    rows = []
    for record in records:
        room = record.room
        sets = sorted(
            [*record.parameter_sets.all(), *record.extended_parameter_sets.all(), *record.parameter_sets_for_storage.all()],
            key=lambda s: s.time or datetime.min.time(),
        )
        base = [
            (_fmt_date(record.created_at), False),
            (room.building.building_number if room and room.building_id else '', False),
            (room.room_number if room else '', False),
            (_room_type(room) if room else '', False),
            (_person(record.responsible), False),
        ]
        instruments = _instruments(record)
        metrics = _room_metrics(room) if room else []

        for parameter_set in sets or [None]:
            cells, violations = [], []
            for key, label, low, high in metrics:
                value = getattr(parameter_set, key, None) if parameter_set else None
                invalid = _out_of_range(value, low, high)
                if invalid:
                    violations.append(label)
                cells.append((_num(value), invalid))
            time_value = parameter_set.time.strftime('%H:%M') if parameter_set and parameter_set.time else ''
            rows.append([*base, (time_value, False), *cells, (', '.join(violations), bool(violations)), (instruments, False)])

    title = 'Параметры по помещениям'
    wb = _build_workbook('Помещения', title, _filters_summary(params, 'rooms'), ROOM_COLUMNS, rows)
    return _response(wb, _filename(title, params))


BUILDING_COLUMNS = [
    ('Дата', 12, None),
    ('Здание', 9, None),
    ('Ответственный', 26, None),
    ('Время', 8, None),
    ('Напряжение сети, В', 13, '0.00'),
    ('Частота тока, Гц', 12, '0.00'),
    ('Форма кривой напряжения', 18, None),
    ('Коэффициент гармоник, %', 13, '0.0'),
    ('Отклонения от нормы', 24, None),
    ('Средства измерений', 48, None),
]


def export_buildings(records, params):
    rows = []
    for record in records:
        building = record.building
        base = [
            (_fmt_date(record.created_at), False),
            (building.building_number if building else '', False),
            (_person(record.responsible), False),
        ]
        instruments = _instruments(record)
        sets = sorted(record.parameter_sets.all(), key=lambda s: s.time or datetime.min.time())

        for parameter_set in sets or [None]:
            voltage = getattr(parameter_set, 'voltage', None)
            frequency = getattr(parameter_set, 'frequency', None)
            waveform = getattr(parameter_set, 'waveform_shape', '') or ''
            harmonic = getattr(parameter_set, 'harmonic_coefficient', None)

            voltage_bad = bool(building) and _out_of_range(voltage, building.voltage_min, building.voltage_max)
            frequency_bad = bool(building) and _out_of_range(frequency, building.frequency_min, building.frequency_max)
            waveform_bad = bool(waveform) and waveform.strip().lower() != EXPECTED_WAVEFORM
            violations = [label for label, bad in (
                ('напряжение', voltage_bad), ('частота', frequency_bad), ('форма кривой', waveform_bad),
            ) if bad]

            rows.append([
                *base,
                (parameter_set.time.strftime('%H:%M') if parameter_set and parameter_set.time else '', False),
                (_num(voltage), voltage_bad),
                (_num(frequency), frequency_bad),
                (waveform, waveform_bad),
                (_num(harmonic), False),
                (', '.join(violations), bool(violations)),
                (instruments, False),
            ])

    title = 'Параметры по зданиям'
    wb = _build_workbook('Здания', title, _filters_summary(params, 'buildings'), BUILDING_COLUMNS, rows)
    return _response(wb, _filename(title, params))
