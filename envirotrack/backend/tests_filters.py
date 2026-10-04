import datetime
from io import BytesIO

from django.contrib.auth.models import User
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from backend.models import (
    Building, BuildingEnviromentalParameters, BuildingParameterSet, EnviromentalParameters,
    ParameterSet, Responsible, Room, UserFilterPreference,
)


class RoomFilterTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('ivanov', password='pass12345')
        cls.other_user = User.objects.create_user('petrova', password='pass12345')
        cls.me = Responsible.objects.create(user=cls.user, last_name='Иванов', first_name='С', patronymic='П')
        cls.other = Responsible.objects.create(user=cls.other_user, last_name='Петрова', first_name='А', patronymic='В')

        b1 = Building.objects.create(building_number='1')
        b7 = Building.objects.create(building_number='7')
        cls.my_room = Room.objects.create(building=b1, room_number='101')
        cls.my_room.responsible_persons.add(cls.me)
        cls.other_room = Room.objects.create(building=b7, room_number='101')
        cls.other_room.responsible_persons.add(cls.other)

        today = datetime.date(2026, 10, 1)
        cls.my_record = EnviromentalParameters.objects.create(room=cls.my_room, responsible=cls.me, created_at=today)
        cls.other_record = EnviromentalParameters.objects.create(
            room=cls.other_room, responsible=cls.other, created_at=today - datetime.timedelta(days=5),
        )
        # Ivanov filled in a record for someone else's room — it still counts as "mine".
        cls.my_record_in_other_room = EnviromentalParameters.objects.create(
            room=cls.other_room, responsible=cls.me, created_at=today,
        )

    def ids(self, response):
        self.assertEqual(response.status_code, 200)
        return {item['id'] for item in response.json()}

    def test_mine_returns_own_rooms_and_own_records(self):
        self.client.force_authenticate(self.user)
        got = self.ids(self.client.get('/api/filterParameters/', {'mine': '1'}))
        self.assertEqual(got, {self.my_record.id, self.my_record_in_other_room.id})

    def test_mine_for_anonymous_is_empty(self):
        self.assertEqual(self.ids(self.client.get('/api/filterParameters/', {'mine': '1'})), set())

    def test_rooms_filter_uses_ids_not_numbers(self):
        got = self.ids(self.client.get('/api/filterParameters/', {'rooms': str(self.other_room.id)}))
        self.assertEqual(got, {self.other_record.id, self.my_record_in_other_room.id})

    def test_open_ended_date_range(self):
        got = self.ids(self.client.get('/api/filterParameters/', {'start_date': '2026-09-30'}))
        self.assertEqual(got, {self.my_record.id, self.my_record_in_other_room.id})

    def test_list_payload_has_building_number(self):
        data = self.client.get('/api/filterParameters/', {'rooms': str(self.my_room.id)}).json()
        self.assertEqual(data[0]['room']['building_number'], '1')


class FilterPreferenceTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('ivanov', password='pass12345')
        cls.other_user = User.objects.create_user('petrova', password='pass12345')

    def test_requires_auth(self):
        self.assertEqual(self.client.get('/api/filter_preferences/rooms/').status_code, 401)

    def test_unknown_scope(self):
        self.client.force_authenticate(self.user)
        self.assertEqual(self.client.get('/api/filter_preferences/nope/').status_code, 404)

    def test_put_sanitizes_and_is_per_user(self):
        self.client.force_authenticate(self.user)
        response = self.client.put('/api/filter_preferences/rooms/', {
            'filters': {'rooms': [3, 'x', 1, 3], 'mine': 1, 'start_date': '2026-10-01', 'end_date': 'oops', 'junk': True},
        }, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['filters'], {
            'responsible': [], 'rooms': [1, 3], 'mine': True, 'start_date': '2026-10-01', 'end_date': '',
        })

        self.client.force_authenticate(self.other_user)
        self.assertIsNone(self.client.get('/api/filter_preferences/rooms/').json()['filters'])

    def test_put_overwrites_and_delete_resets(self):
        self.client.force_authenticate(self.user)
        self.client.put('/api/filter_preferences/buildings/', {'filters': {'buildings': [1]}}, format='json')
        self.client.put('/api/filter_preferences/buildings/', {'filters': {'buildings': [2]}}, format='json')
        self.assertEqual(UserFilterPreference.objects.filter(user=self.user).count(), 1)
        self.assertEqual(self.client.get('/api/filter_preferences/buildings/').json()['filters']['buildings'], [2])

        self.assertEqual(self.client.delete('/api/filter_preferences/buildings/').status_code, 204)
        self.assertIsNone(self.client.get('/api/filter_preferences/buildings/').json()['filters'])


class ExcelExportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('ivanov', password='pass12345')
        cls.me = Responsible.objects.create(user=cls.user, last_name='Иванов', first_name='Сергей', patronymic='П')
        cls.other = Responsible.objects.create(last_name='Петрова', first_name='Анна', patronymic='В')

        cls.b1 = Building.objects.create(building_number='1', voltage_min=207, voltage_max=253)
        cls.b7 = Building.objects.create(building_number='7', voltage_min=207, voltage_max=253)
        cls.room = Room.objects.create(building=cls.b1, room_number='101', humidity_min=30, humidity_max=60)
        cls.room.responsible_persons.add(cls.me)
        cls.other_room = Room.objects.create(building=cls.b7, room_number='101')

        def room_record(room, who, day, humidity):
            record = EnviromentalParameters.objects.create(room=room, responsible=who, created_at=datetime.date(2026, 10, day))
            record.parameter_sets.add(ParameterSet.objects.create(
                temperature_celsius=21, humidity_percentage=humidity, pressure_kpa=100, pressure_mmhg=750, time='09:00',
            ))
            return record

        room_record(cls.room, cls.me, 1, 45.5)
        room_record(cls.room, cls.me, 3, 65.0)
        room_record(cls.other_room, cls.other, 2, 40.0)

        record = BuildingEnviromentalParameters.objects.create(building=cls.b1, responsible=cls.me, created_at=datetime.date(2026, 10, 2))
        record.parameter_sets.add(BuildingParameterSet.objects.create(
            voltage=260, frequency=50, time='10:00', waveform_shape='искажённая', harmonic_coefficient=4.3,
        ))
        BuildingEnviromentalParameters.objects.create(building=cls.b7, responsible=cls.other, created_at=datetime.date(2026, 10, 2))

    def sheet(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        return load_workbook(BytesIO(response.content)).active

    @staticmethod
    def data_rows(ws):
        return [row for row in ws.iter_rows(min_row=6, values_only=True) if any(cell is not None for cell in row)]

    def test_rooms_export_applies_filters_and_marks_violations(self):
        response = self.client.get('/api/export-parameters/', {
            'rooms': str(self.room.id), 'start_date': '2026-10-01', 'end_date': '2026-10-31',
        })
        ws = self.sheet(response)
        rows = self.data_rows(ws)

        self.assertEqual(len(rows), 2)
        self.assertEqual({(r[1], r[2]) for r in rows}, {('1', '101')})
        self.assertIn('с 01.10.2026 по 31.10.2026', ws['A2'].value)
        self.assertIn('101 (зд. 1)', ws['A2'].value)

        newest = rows[0]
        self.assertEqual(newest[0], '03.10.2026')
        self.assertEqual(newest[7], 65.0)
        self.assertEqual(newest[13], 'влажность')
        self.assertEqual(ws.cell(row=6, column=8).font.color.rgb[-6:], 'C62F2F')
        self.assertIn("filename*=UTF-8''", response['Content-Disposition'])

    def test_rooms_export_mine_and_responsible(self):
        self.client.force_authenticate(self.user)
        mine = self.data_rows(self.sheet(self.client.get('/api/export-parameters/', {'mine': '1'})))
        self.assertEqual({r[4] for r in mine}, {'Иванов Сергей П'})

        theirs = self.data_rows(self.sheet(self.client.get('/api/export-parameters/', {'responsible': str(self.other.id)})))
        self.assertEqual([(r[1], r[4]) for r in theirs], [('7', 'Петрова Анна В')])

    def test_single_date_no_longer_crashes(self):
        rows = self.data_rows(self.sheet(self.client.get('/api/export-parameters/', {'date': '2026-10-02'})))
        self.assertEqual(len(rows), 1)

    def test_buildings_export_has_waveform_and_harmonics(self):
        ws = self.sheet(self.client.get('/api/export-parameters-buildings/', {'buildings': str(self.b1.id)}))
        rows = self.data_rows(ws)
        self.assertEqual(len(rows), 1)
        date, building, person, time, voltage, frequency, waveform, harmonic, violations, _ = rows[0]
        self.assertEqual((building, voltage, waveform, harmonic), ('1', 260.0, 'искажённая', 4.3))
        self.assertEqual(violations, 'напряжение, форма кривой')

    def test_record_without_sets_still_listed(self):
        rows = self.data_rows(self.sheet(self.client.get('/api/export-parameters-buildings/', {'buildings': str(self.b7.id)})))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], '7')
